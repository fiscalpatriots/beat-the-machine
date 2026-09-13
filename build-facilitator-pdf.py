# Builds Second-Pass-Facilitator-Guide.pdf, the one page facilitator's guide.
#
# Content of record: FACILITATOR-ONE-PAGE.md. Nothing is typed into this script. The long form
# is FACILITATOR-GUIDE.md and it is not built to PDF, because it does not fit on one page.
#
# Typography and the build route are carried over from
# Coach Dashboard HQ/drafts-2026-09-10-fulbright-packet/build-packet-pdfs.py so the document
# reads as part of the same set: letter, 0.6 inch margins, Calibri, black only, a thin rule
# under every section heading, title centred and bold, body 10.5 point and never smaller.
# python-docx writes the docx, docx2pdf converts it, and PyMuPDF verifies the file on disk:
# page count, the bottom of the last rendered line against the bottom margin, and the metadata
# read back out. FIT is the ladder the builder walks. Spacing and leading give, the body size
# does not.

import os
import re
import sys
import io
from datetime import datetime, timezone

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')

import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE = os.path.dirname(os.path.abspath(__file__))
AUTHOR = 'Khaled Alkurd'
SUBJECT = 'Second Pass: Beat the Machine, running the session in forty minutes'

# source markdown, output basename, maximum pages allowed
JOBS = [
    ('FACILITATOR-ONE-PAGE.md', 'Second-Pass-Facilitator-Guide', 1),
]

FIT = [(10.5, 1.00, 1.00), (10.5, 0.80, 1.00), (10.5, 0.60, 1.00), (10.5, 0.45, 1.00),
       (10.5, 0.45, 0.98), (10.5, 0.35, 0.97), (10.5, 0.25, 0.96)]

TEXT_WIDTH = Inches(7.3)
BULLET_HANG = Inches(0.25)
NUMBER_HANG = Inches(0.30)

S_TITLE_AFTER = 8.0
S_HEAD_BEFORE = 10.0
S_HEAD_AFTER = 4.0
S_PARA_BEFORE = 6.0
S_ITEM_BEFORE = 2.5
S_TABLE_BEFORE = 5.0
S_TABLE_AFTER = 6.0

BAD_CHARS = {'—': 'em dash', '–': 'en dash', '‘': 'left single quote',
             '’': 'right single quote', '“': 'left double quote',
             '”': 'right double quote', '…': 'ellipsis'}

SMALL_WORDS = {'a', 'an', 'and', 'as', 'at', 'but', 'by', 'for', 'from', 'in', 'nor', 'of',
               'on', 'or', 'the', 'to', 'with'}


# ---------------------------------------------------------------------------
# 1. Read the content of record
# ---------------------------------------------------------------------------

def title_case(text):
    words = text.split(' ')
    out = []
    for i, w in enumerate(words):
        bare = w.strip(',.;:').lower()
        if i not in (0, len(words) - 1) and bare in SMALL_WORDS:
            out.append(w[0].lower() + w[1:] if w else w)
        elif w[:1].islower():
            out.append(w[0].upper() + w[1:])
        else:
            out.append(w)
    return ' '.join(out)


def parse(path):
    """Markdown to a flat list of blocks: title, heading, para, bullet, number, table."""
    raw = io.open(path, encoding='utf-8').read().splitlines()
    blocks = []
    i = 0
    pending = []
    item = None

    def flush_para():
        if pending:
            blocks.append(('para', ' '.join(pending), None))
            del pending[:]

    def flush_item():
        if item is not None:
            blocks.append((item[0], ' '.join(item[2]), item[1]))

    while i < len(raw):
        line = raw[i].rstrip()
        stripped = line.strip()

        if not stripped:
            flush_item()
            item = None
            flush_para()
            i += 1
            continue

        if stripped == '---':
            flush_item(); item = None; flush_para()
            i += 1
            continue

        if line.startswith('# '):
            flush_item(); item = None; flush_para()
            blocks.append(('title', title_case(line[2:].strip()), None))
            i += 1
            continue

        if line.startswith('## '):
            flush_item(); item = None; flush_para()
            blocks.append(('heading', line[3:].strip(), None))
            i += 1
            continue

        if stripped.startswith('|'):
            flush_item(); item = None; flush_para()
            rows = []
            while i < len(raw) and raw[i].strip().startswith('|'):
                cells = [c.strip() for c in raw[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r':?-{2,}:?', c) for c in cells):
                    rows.append(cells)
                i += 1
            blocks.append(('table', rows, None))
            continue

        m = re.match(r'^(\s*)-\s+(.*)$', line)
        if m and len(m.group(1)) < 2:
            flush_item(); flush_para()
            item = ('bullet', None, [m.group(2).strip()])
            i += 1
            continue

        m = re.match(r'^(\s*)(\d+)\.\s+(.*)$', line)
        if m and len(m.group(1)) < 2:
            flush_item(); flush_para()
            item = ('number', m.group(2) + '.', [m.group(3).strip()])
            i += 1
            continue

        if item is not None and line[:1] == ' ':
            item[2].append(stripped)
            i += 1
            continue

        flush_item(); item = None
        pending.append(stripped)
        i += 1

    flush_item()
    flush_para()
    return blocks


# ---------------------------------------------------------------------------
# 2. Build the document
# ---------------------------------------------------------------------------

def rule(paragraph, sz='6', color='000000'):
    pPr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), sz)
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), color)
    borders.append(bottom)
    pPr.append(borders)


def keep_with_next(paragraph):
    pPr = paragraph._p.get_or_add_pPr()
    pPr.append(OxmlElement('w:keepNext'))


def emit_inline(paragraph, text, size, base_bold=False):
    """Renders **bold** and strips the backticks off code spans."""
    text = text.replace('`', '')
    for piece in re.split(r'(\*\*.+?\*\*)', text):
        if not piece:
            continue
        bold = base_bold
        if piece.startswith('**') and piece.endswith('**') and len(piece) > 4:
            piece = piece[2:-2]
            bold = True
        r = paragraph.add_run(piece)
        r.bold = bold
        r.font.size = Pt(size)
        r.font.name = 'Calibri'


def cell_borders(cell, color='9A9A9A', sz='4'):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement('w:tcBorders')
    for edge in ('top', 'left', 'bottom', 'right'):
        e = OxmlElement('w:' + edge)
        e.set(qn('w:val'), 'single')
        e.set(qn('w:sz'), sz)
        e.set(qn('w:space'), '0')
        e.set(qn('w:color'), color)
        borders.append(e)
    tcPr.append(borders)


def mark_size(paragraph, half_points):
    pPr = paragraph._p.get_or_add_pPr()
    rPr = pPr.find(qn('w:rPr'))
    if rPr is None:
        rPr = OxmlElement('w:rPr')
        pPr.append(rPr)
    for tag in ('w:sz', 'w:szCs'):
        sz = OxmlElement(tag)
        sz.set(qn('w:val'), str(half_points))
        rPr.append(sz)


def build(blocks, out_path, doc_title, body, space, leading):
    d = docx.Document()
    for s in d.sections:
        s.top_margin = s.bottom_margin = s.left_margin = s.right_margin = Inches(0.6)
    n = d.styles['Normal']
    n.font.name = 'Calibri'
    n.font.size = Pt(body)
    n.font.color.rgb = RGBColor(0, 0, 0)
    n.paragraph_format.space_after = Pt(0)
    n.paragraph_format.space_before = Pt(0)
    n.paragraph_format.line_spacing = leading

    def para(before=0.0, after=0.0, hang=None, spacer=False):
        p = d.add_paragraph()
        pf = p.paragraph_format
        pf.space_before = Pt(before * space)
        pf.space_after = Pt(after * space)
        pf.line_spacing = leading
        if hang is not None:
            pf.left_indent = hang
            pf.first_line_indent = -hang
            pf.tab_stops.add_tab_stop(hang, WD_TAB_ALIGNMENT.LEFT)
        if spacer:
            mark_size(p, 2)
        return p

    first_para_after_head = False
    for kind, payload, marker in blocks:
        if kind == 'title':
            p = para(after=S_TITLE_AFTER)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(payload)
            r.bold = True
            r.font.size = Pt(15)
            r.font.name = 'Calibri'
            keep_with_next(p)

        elif kind == 'heading':
            p = para(before=S_HEAD_BEFORE, after=S_HEAD_AFTER)
            r = p.add_run(payload)
            r.bold = True
            r.font.size = Pt(body)
            r.font.name = 'Calibri'
            rPr = r._element.get_or_add_rPr()
            sp = OxmlElement('w:spacing')
            sp.set(qn('w:val'), '20')
            rPr.append(sp)
            rule(p)
            keep_with_next(p)
            first_para_after_head = True
            continue

        elif kind == 'para':
            p = para(before=0.0 if first_para_after_head else S_PARA_BEFORE)
            emit_inline(p, payload, body)

        elif kind in ('bullet', 'number'):
            hang = BULLET_HANG if kind == 'bullet' else NUMBER_HANG
            p = para(before=S_ITEM_BEFORE, hang=hang)
            lead = p.add_run((u'•' if kind == 'bullet' else marker) + '\t')
            lead.font.size = Pt(body)
            lead.font.name = 'Calibri'
            emit_inline(p, payload, body)

        elif kind == 'table':
            rows = payload
            cols = max(len(r) for r in rows)
            t = d.add_table(rows=0, cols=cols)
            t.alignment = WD_TABLE_ALIGNMENT.LEFT
            t.autofit = False
            if cols == 2:
                widths = [Inches(1.45), TEXT_WIDTH - Inches(1.45)]
            elif cols == 3:
                # A first column of short values, such as a minutes count, does not need the
                # width a question column does, so the row is measured before it is set.
                first_max = max(len(r[0]) for r in rows if r)
                if first_max <= 12:
                    widths = [Inches(0.8), Inches(1.3), Inches(5.2)]
                else:
                    widths = [Inches(2.3), Inches(1.5), Inches(3.5)]
            else:
                widths = [Inches(7.3 / cols)] * cols
            for ri, cells in enumerate(rows):
                row = t.add_row()
                for ci in range(cols):
                    cell = row.cells[ci]
                    cell.width = widths[ci]
                    cp = cell.paragraphs[0]
                    cp.paragraph_format.space_before = Pt(0.8 * space)
                    cp.paragraph_format.space_after = Pt(0.8 * space)
                    cp.paragraph_format.line_spacing = leading
                    txt = cells[ci] if ci < len(cells) else ''
                    emit_inline(cp, txt, body, base_bold=(ri == 0))
                    cell_borders(cell)
            t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
            t._tbl.addprevious(para(before=S_TABLE_BEFORE, spacer=True)._p)
            para(before=S_TABLE_AFTER, spacer=True)

        first_para_after_head = False

    cp = d.core_properties
    cp.author = AUTHOR
    cp.last_modified_by = AUTHOR
    cp.title = doc_title
    cp.subject = SUBJECT
    now = datetime.now(timezone.utc).replace(microsecond=0)
    cp.created = now
    cp.modified = now
    cp.revision = 1
    d.save(out_path)
    return out_path


# ---------------------------------------------------------------------------
# 3. Convert and verify by rendering
# ---------------------------------------------------------------------------

def to_pdf(docx_path):
    from docx2pdf import convert
    pdf_path = docx_path[:-5] + '.pdf'
    if os.path.exists(pdf_path):
        os.remove(pdf_path)
    convert(docx_path, pdf_path)
    return pdf_path


def page_count(pdf_path):
    import fitz
    doc = fitz.open(pdf_path)
    k = doc.page_count
    doc.close()
    return k


def stamp_and_verify(pdf_path, doc_title, png_dir):
    import fitz
    doc = fitz.open(pdf_path)
    meta = doc.metadata or {}
    if (meta.get('title') or '') != doc_title or (meta.get('author') or '') != AUTHOR:
        doc.set_metadata({'title': doc_title, 'author': AUTHOR, 'subject': SUBJECT,
                          'creator': 'Microsoft Word', 'producer': 'docx2pdf'})
        doc.saveIncr()
    doc.close()

    doc = fitz.open(pdf_path)
    meta = doc.metadata or {}
    report = {'pages': doc.page_count, 'title': meta.get('title'), 'author': meta.get('author'),
              'clipped': [], 'chars': 0}
    for pno in range(doc.page_count):
        page = doc[pno]
        rect = page.rect
        report['chars'] += len(page.get_text())
        for b in page.get_text('dict')['blocks']:
            if b['type'] != 0:
                continue
            for l in b['lines']:
                x0, y0, x1, y1 = l['bbox']
                if x0 < 1 or y0 < 1 or x1 > rect.width - 1 or y1 > rect.height - 1:
                    report['clipped'].append((pno + 1, [round(v, 1) for v in l['bbox']]))
        if pno == 0:
            pix = page.get_pixmap(dpi=110)
            png = os.path.join(png_dir, os.path.basename(pdf_path)[:-4] + '-p1.png')
            pix.save(png)
            report['png'] = png
            lines = [l for b in page.get_text('dict')['blocks'] if b['type'] == 0
                     for l in b['lines']]
            report['last_baseline'] = round(max(l['bbox'][3] for l in lines), 1)
            report['page_h'] = round(rect.height, 1)
            report['bottom_margin_at'] = round(rect.height - 43.2, 1)
    doc.close()
    return report


def banned(blocks):
    text = []
    for kind, payload, _ in blocks:
        if kind == 'table':
            text.extend(c for row in payload for c in row)
        else:
            text.append(payload)
    joined = ' '.join(text)
    return sorted({label for ch, label in BAD_CHARS.items() if ch in joined})


def run(source, basename, max_pages):
    path = os.path.join(HERE, source)
    blocks = parse(path)
    docx_path = os.path.join(HERE, basename + '.docx')
    doc_title = next(p for k, p, _ in blocks if k == 'title')

    chosen = None
    for body, space, leading in FIT:
        build(blocks, docx_path, doc_title, body, space, leading)
        pdf_path = to_pdf(docx_path)
        pages = page_count(pdf_path)
        print('  try %.1f pt / spacing %.2f / leading %.2f -> %d page(s)'
              % (body, space, leading, pages))
        if pages <= max_pages:
            chosen = (body, space, leading)
            break
    if chosen is None:
        chosen = FIT[-1]

    rep = stamp_and_verify(pdf_path, doc_title, HERE)
    print('  source             : %s (%d blocks)' % (source, len(blocks)))
    print('  heading / pdf title: %r' % doc_title)
    print('  typography         : Calibri %.1f pt, leading %.2f, spacing scale %.2f, '
          '0.6 inch margins' % (chosen[0], chosen[2], chosen[1]))
    print('  pdf pages          : %d of %d allowed  %s'
          % (rep['pages'], max_pages, 'PASS' if rep['pages'] <= max_pages else 'FAIL'))
    print('  page 1 last line   : %.1f pt of %.1f, bottom margin starts at %.1f'
          % (rep['last_baseline'], rep['page_h'], rep['bottom_margin_at']))
    print('  text outside page  : %s' % (rep['clipped'] if rep['clipped'] else 'none, PASS'))
    print('  extracted chars    : %d' % rep['chars'])
    print('  pdf metadata       : author=%r title=%r' % (rep['author'], rep['title']))
    print('  page 1 image       : %s' % os.path.basename(rep['png']))
    print('  banned characters  : %s' % (', '.join(banned(blocks)) or 'none'))
    print('  docx               : %s' % os.path.basename(docx_path))
    print('')
    return rep['pages'] <= max_pages and not rep['clipped']


if __name__ == '__main__':
    want = sys.argv[1:]
    jobs = [j for j in JOBS if not want or any(w.lower() in j[1].lower() for w in want)]
    ok = True
    for source, basename, max_pages in jobs:
        print(basename)
        ok = run(source, basename, max_pages) and ok
    print('all %d within their page allowance and nothing clipped: %s' % (len(jobs), ok))
